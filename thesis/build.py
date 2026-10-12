"""Build the evidence-backed thesis with native Word equations and editable tables."""
import json
from pathlib import Path
import re
import statistics

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Cm
from PIL import Image, ImageDraw, ImageFont

ROOT=Path(__file__).resolve().parent.parent
OUT=ROOT/'thesis'
TITLE='受限环境下的语义自适应视频传输系统设计与实现'
NAMES={'full':'全帧 Q85','uniform':'全局自适应','static-roi':'固定区域','semantic':'动态优先'}


def xml(tag,**attrs):
    element=OxmlElement(tag)
    for name,value in attrs.items(): element.set(qn(name),str(value))
    return element


def math_run(text):
    run=xml('m:r');node=xml('m:t');node.text=str(text);run.append(node);return run


def sub(base,index):
    node=xml('m:sSub');body=xml('m:e');body.append(math_run(base));tail=xml('m:sub');tail.append(math_run(index));node.extend([body,tail]);return node


def fraction(top,bottom):
    node=xml('m:f');num=xml('m:num');den=xml('m:den')
    for part in top: num.append(math_run(part) if isinstance(part,str) else part)
    for part in bottom: den.append(math_run(part) if isinstance(part,str) else part)
    node.extend([num,den]);return node


def equations():
    return {
        'latency':('2-1',[sub('L','k'),' = ',sub('r','k'),' + θ − ',sub('s','k')]),
        'age':('2-2',['Δ(t) = t + θ − ',sub('s','latest')]),
        'psnr':('2-3',['PSNR = 10 log₁₀ ',fraction(['255²'],['MSE'])]),
        'priority':('4-1',[sub('p','i'),' = ',sub('w','i'),' [1 + min(12',sub('c','i'),', 2)] + min(',fraction([sub('a','i')],['250 ms']),', 12)']),
        'budget':('4-2',[sub('B','max'),' = ',fraction(['1000 R'],['8 f'])]),
        'pacing':('4-3',[sub('T','send'),' = max(',fraction(['1000'],['f']),', ',fraction(['8 S'],['R']),') ms']),
        'offset':('5-1',['θ = ',fraction(['(',sub('t','1'),' − ',sub('t','0'),') + (',sub('t','2'),' − ',sub('t','3'),')'],['2'])]),
        'rtt':('5-2',['RTT = (',sub('t','3'),' − ',sub('t','0'),') − (',sub('t','2'),' − ',sub('t','1'),')'])
    }


def set_font(run,size=12,bold=False):
    run.font.name='Times New Roman';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun')


class Writer:
    def __init__(self):
        self.doc=Document();section=self.doc.sections[0]
        section.page_width=Inches(8.5);section.page_height=Inches(11)
        section.top_margin=section.bottom_margin=Inches(.85);section.left_margin=section.right_margin=Inches(1)
        section.header_distance=section.footer_distance=Inches(.4)
        for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Caption']:
            style=self.doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
            style.font.size=Pt(12);style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'SimSun')
            borders=style.element.find('.//'+qn('w:pBdr'))
            if borders is not None: borders.getparent().remove(borders)
        normal=self.doc.styles['Normal'].paragraph_format
        normal.line_spacing=1.5;normal.space_after=Pt(6);normal.widow_control=True
        for level,size in [(1,16),(2,13)]:
            style=self.doc.styles[f'Heading {level}'];style.font.size=Pt(size);style.font.bold=True
            style.paragraph_format.space_before=Pt(12);style.paragraph_format.space_after=Pt(8)
            style.paragraph_format.keep_with_next=True
        self.doc.styles['Title'].font.size=Pt(23)
        section.different_first_page_header_footer=True
        header=section.header.paragraphs[0];header.alignment=WD_ALIGN_PARAGRAPH.CENTER
        set_font(header.add_run('语义自适应视频传输系统    本科论文初稿'),9)
        footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
        field=xml('w:fldSimple',**{'w:instr':'PAGE'});footer._p.append(field)

    def p(self,text,indent=True):
        paragraph=self.doc.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent: paragraph.paragraph_format.first_line_indent=Pt(24)
        set_font(paragraph.add_run(text));return paragraph

    def caption(self,text):
        paragraph=self.doc.add_paragraph(style='Caption');paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before=Pt(6);paragraph.paragraph_format.space_after=Pt(8)
        set_font(paragraph.add_run(text),10.5);return paragraph

    def table(self,title,headers,rows,widths):
        self.caption(title).paragraph_format.keep_with_next=True
        table=self.doc.add_table(rows=1,cols=len(headers));table.autofit=False;table.alignment=WD_TABLE_ALIGNMENT.CENTER
        borders=xml('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            borders.append(xml('w:'+edge,**{'w:val':'single','w:sz':4,'w:color':'D9D9D9'}))
        table._tbl.tblPr.append(borders)
        margins=xml('w:tblCellMar')
        for edge in ['top','bottom','left','right']: margins.append(xml('w:'+edge,**{'w:w':85,'w:type':'dxa'}))
        table._tbl.tblPr.append(margins)
        table.rows[0]._tr.get_or_add_trPr().append(xml('w:tblHeader'))
        for row_index,values in enumerate([headers]+rows):
            cells=table.rows[0].cells if row_index==0 else table.add_row().cells
            for index,(cell,value) in enumerate(zip(cells,values)):
                cell.width=Cm(widths[index]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                fill='E5E5E5' if row_index==0 else ('F7F7F7' if row_index%2==0 else 'FFFFFF')
                cell._tc.get_or_add_tcPr().append(xml('w:shd',**{'w:fill':fill}))
                paragraph=cell.paragraphs[0];paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER if len(str(value))<25 else WD_ALIGN_PARAGRAPH.LEFT
                paragraph.paragraph_format.line_spacing=1.15;paragraph.paragraph_format.space_after=Pt(0)
                set_font(paragraph.add_run(str(value)),10.5,row_index==0)
            table.rows[row_index]._tr.get_or_add_trPr().append(xml('w:cantSplit'))
        for index,column in enumerate(table.columns): column.width=Cm(widths[index])
        self.doc.add_paragraph().paragraph_format.space_after=Pt(0)

    def equation(self,key):
        label,parts=equations()[key];paragraph=self.doc.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before=Pt(5);paragraph.paragraph_format.space_after=Pt(9)
        expression=xml('m:oMath')
        for part in parts: expression.append(math_run(part) if isinstance(part,str) else part)
        paragraph._p.append(expression);set_font(paragraph.add_run('    ('+label+')'),11)


def architecture():
    image=Image.new('RGB',(1600,500),'white');draw=ImageDraw.Draw(image)
    font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',36)
    small=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',32)
    labels=['最新输入','任务区域与预算','同帧背景与重点块','解码并一次合成']
    for i,label in enumerate(labels):
        x=35+i*397;draw.rectangle((x,90,x+335,205),outline='#333333',width=3)
        draw.text((x+167,147),label,font=font,fill='black',anchor='mm')
        if i<3:
            draw.line((x+335,147,x+390,147),fill='black',width=3);draw.polygon([(x+390,147),(x+377,140),(x+377,154)],fill='black')
    draw.line([(1392,205),(1392,285),(200,285),(200,205)],fill='#555555',width=3)
    draw.polygon([(200,205),(193,220),(207,220)],fill='#555555')
    draw.rectangle((495,258,1115,312),fill='white');draw.text((805,285),'显示完成 ACK   最多一帧在途',font=small,anchor='mm',fill='black')
    draw.text((800,360),'独立校时连接   四时间戳   延迟与当前帧龄',font=small,anchor='mm',fill='black')
    draw.text((800,420),'旧系统继续运行   新容器端口 8088   不含设备动作接口',font=small,anchor='mm',fill='#444444')
    destination=OUT/'figures/architecture.png';destination.parent.mkdir(parents=True,exist_ok=True);image.save(destination);return destination


def freshness_chart(groups):
    image=Image.new('RGB',(1400,1100),'white');draw=ImageDraw.Draw(image)
    font=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',32)
    title=ImageFont.truetype('C:/Windows/Fonts/msyh.ttc',42)
    draw.text((25,20),'新鲜帧时间覆盖率',font=title,fill='black')
    draw.text((25,80),'柱为四次均值，线为观测最小值至最大值，非置信区间',font=font,fill='#444444')
    for tick in (0,25,50,75,100):
        x=400+8*tick;draw.line((x,160,x,995),fill='#dddddd',width=2)
        draw.text((x,1025),str(tick)+'%',font=font,fill='black',anchor='mm')
    for band,kbps in enumerate((512,1024,2048)):
        top=180+band*275
        draw.text((25,top),str(kbps)+' kbps',font=font,fill='black')
        for index,mode in enumerate(['full','uniform','static-roi','semantic']):
            row=next(r for r in groups if r['kbps']==kbps and r['policy']==mode);y=top+index*55
            draw.text((360,y),NAMES[mode],font=font,fill='black',anchor='rm')
            end=400+800*row['fresh_fraction']
            if end>400: draw.rectangle((400,y-13,end,y+13),fill='#457878')
            low,high=400+800*row['fresh_min'],400+800*row['fresh_max']
            draw.line((low,y,high,y),fill='#202020',width=3)
            for x in (low,high):draw.line((x,y-20,x,y+20),fill='#202020',width=3)
            draw.text((1240,y),f'{100*row["fresh_fraction"]:.1f}%',font=font,fill='black',anchor='lm')
    destination=OUT/'figures/freshness.png';destination.parent.mkdir(parents=True,exist_ok=True);image.save(destination);return destination


def main():
    analysis=json.loads((ROOT/'results/analysis.json').read_text(encoding='utf-8'))
    validation=json.loads((ROOT/'results/validation.json').read_text(encoding='utf-8'))
    writer=Writer();doc=writer.doc
    doc.core_properties.title=TITLE;doc.core_properties.subject='测控技术与仪器本科论文初稿';doc.core_properties.author=''
    paragraph=doc.add_paragraph(style='Title');paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before=Pt(80);paragraph.paragraph_format.space_after=Pt(26)
    run=paragraph.add_run(TITLE.replace('自适应视频','自适应\n视频'));set_font(run,23,True)
    for text in ['本科毕业论文初稿','测控技术与仪器','面向任务区域的资源分配与信息新鲜度实验','实验和文档日期 2026年10月12日']:
        p=writer.p(text,False);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(18)
    writer.p('本文包含可运行原型、合成输入与树莓派实测结果。尚未获得真实产线任务真值，未实现工业设备控制。作者、学校及导师信息须由学生依据学校要求补充。',False)
    doc.add_page_break();doc.add_heading('摘要',level=1)
    writer.p('针对带宽和边缘算力受限时整幅视频质量与远程观察新鲜度难以同时维持的问题，本文设计并实现一个独立的语义自适应视频传输原型。语义被限定为任务区域与重要性，而非自动场景理解。系统利用区域权重、局部变化和更新等待时间进行资源分配，以同帧低分辨率背景和重点 JPEG 块构成自包含消息，通过单帧在途确认机制抑制应用层积压，并利用四时间戳校时显示软件段延迟与当前帧龄。')
    writer.p(f'实验包括树莓派离线配对编码 1080 个样本，以及 {analysis["network_sessions"]} 个受限链路主对照单元，共 {analysis["network_seconds"]:.0f} s 正式观察窗口。此外进行了 8 个真实输入对照单元和 28 个反馈消融单元，共核对 84 个网络测量单元。除成功帧时延外，结果同时报告背景与区域质量、合成任务标记保真度及新鲜任务覆盖率。结果表明，区域资源分配可改变质量与新鲜度的取舍，但动态策略不保证在所有条件下优于固定权重；无线抖动和反馈退化会使较低的成功帧延迟与较低的任务覆盖率并存。')
    writer.p('本文将版本冻结、只读输入边界、可追溯实验和不确定性分析纳入系统实现。结论限于已测试的场景与配置，不等价于光学端到端时延，不证明生产线效率提升，也未与硬件 H.264 或所有通信协议进行对比。该原型为设备远程观察和测控链路验证提供可复现实验基础。')
    writer.p('关键词：任务区域；自适应传输；信息新鲜度；树莓派；测量不确定性',False)
    doc.add_page_break();doc.add_heading('Abstract',level=1)
    writer.p('This thesis implements an isolated task-oriented adaptive video prototype for constrained bandwidth and edge computation. Semantics are explicitly defined as operator-selected regions and priorities, not automatic scene understanding. A self-contained frame combines a low-resolution context image with same-frame JPEG region patches. One-frame acknowledgement flow control bounds application backlog, while a separate four-timestamp exchange estimates software latency and information age.',False)
    writer.p(f'Evaluation includes 1,080 paired encoding samples, {analysis["network_sessions"]} primary Raspberry Pi network sessions, eight real-input comparison sessions and 28 feedback-ablation sessions. The study measures regional and background fidelity, synthetic marker preservation, received-frame latency and time-based freshness. Dynamic prioritization does not consistently dominate static regions. Wireless stalls and aggressive budget reduction can reduce task usefulness even when the latency of received frames remains small. Results are restricted to the tested prototype, and do not establish glass-to-glass latency, production benefits or superiority over hardware video codecs.',False)
    writer.p('Keywords: task region; adaptive streaming; information age; Raspberry Pi; measurement uncertainty',False)
    doc.add_page_break();toc_heading=doc.add_paragraph()
    set_font(toc_heading.add_run('目录'),16,True)
    toc_heading.paragraph_format.space_after=Pt(12)
    p=doc.add_paragraph();field=xml('w:fldSimple',**{'w:instr':'TOC \\o "1-2" \\h \\z \\u'});p._p.append(field)
    doc.add_page_break()
    tables={
        'constraints':('表 2-1 原型的输入与资源边界',['项目','配置或限制'],[['画布','1280×720'],['区域','1 至 4 个，总面积不超过 50%'],['有效载荷参数','128 至 12000 kbps'],['目标帧率参数','2 至 20 fps'],['并发视频','最多 2 个独立会话'],['容器','非 root、只读、1 CPU、64 进程'],['控制范围','无执行器动作接口']],[4,12.5]),
        'modules':('表 3-1 模块职责与证据',['模块','职责'],[['source','合成、只读共享输入与独立 USB 适配'],['model 和 encoder','区域校验、预算反馈、编码档位和优先级'],['protocol','自包含二进制格式与重建校验'],['server','会话、时钟、确认与超时管理'],['web','一次合成、区域参数、指标与数据导出'],['experiments','配对实验、真实 TCP 中继和独立复核']],[4,12.5]),
        'policies':('表 4-1 对照策略',['策略','空间质量分配','调度依据'],[['全帧 Q85','完整 720p','不改变编码档位'],['全局自适应','全景统一降级','预算与档位'],['固定区域','低质量背景与重点块','预设区域权重'],['动态优先','低质量背景与重点块','权重、变化、等待时间']],[3.2,6.3,7])
    }
    def result_block(key):
        if key=='validation':
            writer.p(f'本地 Python 测试 {validation["python_tests"]} 项、JavaScript 协议与时钟测试 {validation["javascript_tests"]} 项通过；树莓派容器内 Python 测试 {validation["pi_python_tests"]} 项通过。浏览器功能与布局检查为 {validation["browser_checks"]} 项，覆盖四种视口及四种传输策略。测试结果与部署属性记录于 results/validation.json。')
        elif key=='rd':
            rows=[[r['kbps'],NAMES[r['policy']],f'{r["bytes"]/1000:.1f}',f'{r["weighted_roi_psnr_db"]:.2f}',f'{r["background_psnr_db"]:.2f}',f'{100*r["marker_fraction"]:.1f}'] for r in analysis['rd']]
            writer.table('表 6-1 树莓派配对结果 每单元汇总三场景共 90 帧',['上限 kbps','策略','包 KB','区域 dB','背景 dB','任务 %'],rows,[2,3,2,3,3,3.5])
            writer.p('表中区域 dB 是按权重平均的区域 PSNR，不是整幅图像 PSNR。任务百分比是加权色标保真度。每个编码样本都使用相同的原图参照；预算为名义 10 fps 单帧预算，表格本身不包含网络停顿。')
            subrows=[[r['kbps'],NAMES[r['policy']],f'{100*r["region_success"]["process-A"]:.1f}',f'{100*r["region_success"]["process-B"]:.1f}',f'{r["encode_ms"]:.2f}'] for r in analysis['rd'] if r['policy'] in ('static-roi','semantic')]
            writer.table('表 6-2 区域分配差异和编码开销',['上限 kbps','策略','A 成功 %','B 成功 %','编码 ms'],subrows,[2.5,4,3.5,3.5,3])
        elif key=='network':
            writer.p(f'正式结果汇总两轮、两种纹理场景，每种带宽与策略共 4 个测量单元。总计 {analysis["network_sessions"]} 单元、{analysis["network_seconds"]:.0f} s 正式窗口，其中 {analysis["failed_sessions"]} 单元出现连接关闭或超时记录，{analysis["clock_warning_sessions"]} 单元触发校时质量警告。没有剔除这些失败单元。两轮是重复筛查，样本量仍不足以作强统计推断。')
            rows=[]
            for r in analysis['network']:
                rows.append([r['kbps'],NAMES[r['policy']],r['received_frames'],f'{r["latency_p50_ms"]:.1f}' if r['latency_p50_ms'] is not None else '无帧',f'{r["latency_p95_ms"]:.1f}' if r['latency_p95_ms'] is not None else '无帧',f'{100*r["fresh_fraction"]:.1f}',f'{100*r["task_fraction"]:.1f}'])
            writer.table('表 6-3 树莓派受限链路 成功帧时延与时间覆盖率',['kbps','策略','帧数','P50 ms','P95 ms','新鲜 %','任务 %'],rows,[1.6,2.7,1.6,2.8,2.8,2.5,2.5])
            writer.p('P50 与 P95 由各组正式窗口内的全部成功帧合并计算；帧数越少，分位数越不稳定。覆盖率的分母包括关闭后的剩余观察时间。新鲜度阈值为 250 ms，任务列还要求合成标记保真。它们不能与无相同任务定义的外部论文百分比直接比较。')
            rows=[[r['kbps'],NAMES[r['policy']],f'{100*r["fresh_min"]:.1f} 至 {100*r["fresh_max"]:.1f}',f'{100*r["conservative_fresh_fraction"]:.1f}',r['clock_warning_sessions']] for r in analysis['network']]
            writer.table('表 6-4 波动范围与校时敏感性 非置信区间',['kbps','策略','新鲜率范围 %','保守新鲜 %','校时警告'],rows,[1.7,3,5.3,3.5,3])
            doc.add_picture(str(freshness_chart(analysis['network'])),width=Inches(6.1))
            doc.paragraphs[-1].paragraph_format.keep_with_next=True
            writer.caption('图 6-1 受限链路新鲜时间覆盖率 250 ms 门限 每组四个单元')
        elif key=='live':
            writer.p('真实摄像头通过只读共享内存接入。此次验收没有停止旧 USB 采集器，也没有为新容器授予 USB 设备权限。下表仅描述本次摄像头场景中的软件发布至 Python 解码链路，不代表工艺识别准确率。')
            rows=[[r['kbps'],NAMES[r['policy']],r['frames'],f'{r["fps"]:.2f}',f'{r["latency_p50_ms"]:.1f}' if r['latency_p50_ms'] is not None else '无帧',f'{100*r["fresh_time_fraction_250ms"]:.1f}'] for r in analysis['live']]
            writer.table('表 6-5 真实摄像头只读链路',['kbps','策略','帧数','fps','P50 ms','新鲜 %'],rows,[2,3.5,2,2.5,3.5,2.5])
        elif key=='feedback':
            fixture=analysis['feedback_ablations']['feedback-pi-fixture.json']
            live=analysis['feedback_ablations']['feedback-pi-live.json']
            writer.p(f'抗抖动反馈补充实验包含合成输入 {fixture["sessions"]} 个测量单元，以及真实摄像头 {live["sessions"]} 个测量单元。单次确认与五样本中值两种模式交替测试；编码策略均为动态区域优先。下表不与原始四策略数据混合求总体分位数。')
            rows=[]
            for r in fixture['groups']:
                rows.append([r['kbps'],'单次确认' if r['feedback']=='single' else '五样本中值',f'{r["latency_p50_ms"]:.1f}' if r['latency_p50_ms'] is not None else '无帧',
                             f'{r["latency_p95_ms"]:.1f}' if r['latency_p95_ms'] is not None else '无帧',f'{100*r["fresh_fraction"]:.1f}',f'{100*r["task_fraction"]:.1f}',f'{100*r["patch_frame_fraction"]:.1f}'])
            writer.table('表 6-6 反馈消融 每档两场景各重复两次',['kbps','反馈','P50 ms','P95 ms','新鲜 %','任务 %','重点帧 %'],rows,[1.5,3.3,2.5,2.5,2.1,2.1,2.5])
            comparisons=[]
            for rate in (512,1024,2048):
                single=next(r for r in fixture['groups'] if r['kbps']==rate and r['feedback']=='single')
                median=next(r for r in fixture['groups'] if r['kbps']==rate and r['feedback']=='median')
                comparisons.append(f'{rate} kbps 下，新鲜任务覆盖率从 {100*single["task_fraction"]:.1f}% 变为 {100*median["task_fraction"]:.1f}%，新鲜时间覆盖率从 {100*single["fresh_fraction"]:.1f}% 变为 {100*median["fresh_fraction"]:.1f}%')
            writer.p('本次观测中，'+'；'.join(comparisons)+'。因此，中值反馈并非每档任务指标都更好，也不能将合成输入的优势直接推广至真实摄像头。')
            rows=[]
            for r in live['groups']:
                rows.append(['单次确认' if r['feedback']=='single' else '五样本中值',f'{r["mean_bytes"]/1000:.2f}',f'{100*r["patch_frame_fraction"]:.1f}',
                             f'{r["latency_p50_ms"]:.1f}' if r['latency_p50_ms'] is not None else '无帧',f'{100*r["fresh_fraction"]:.1f}',r['clock_warnings']])
            writer.table('表 6-7 真实输入反馈对照 1200 kbps 每模式重复两次',['反馈','包 KB','重点帧 %','P50 ms','新鲜 %','校时警告'],rows,[3.5,2.4,2.8,2.8,2.5,2.5])
            writer.p('本次真实输入的中值模式新鲜覆盖率和重点帧比例均低于单次反馈，且有一个单元出现校时质量警告。测量采用交替顺序但并非同时承受同一无线条件，不能把全部差异归因于滤波器。当前证据不足以支持全面替换，因此最终网页和服务器默认保持单次反馈，中值模式仅作为可选实验配置；两种模式均不构成工业性能保证。')
            writer.p('重点帧比例表示正式窗口内至少带一个高质量区域块的成功帧占比，分母为成功收到的帧，不是全部观察时间。包长增加和该比例提高都不能替代任务真值。消融结果文件的完整 SHA256 见附录。')
    content=(OUT/'manuscript.md').read_text(encoding='utf-8')
    for block in content.split('\n\n'):
        block=block.strip()
        if not block: continue
        if block.startswith('# '):
            if block!='# 1 绪论': doc.add_page_break()
            doc.add_heading(block[2:],level=1)
        elif block.startswith('## '): doc.add_heading(block[3:],level=2)
        elif block.startswith('{{table:'): writer.table(*tables[block[8:-2]])
        elif block.startswith('{{equation:'): writer.equation(block[11:-2])
        elif block.startswith('{{figure:'):
            doc.add_picture(str(architecture()),width=Inches(6.4));doc.paragraphs[-1].paragraph_format.keep_with_next=True;writer.caption('图 3-1 独立传输系统数据流与确认路径')
        elif block.startswith('{{results:'): result_block(block[10:-2])
        else:
            p=writer.p(block,not block.startswith('['))
            if block.startswith('['):p.alignment=WD_ALIGN_PARAGRAPH.LEFT
    doc.add_page_break();doc.add_heading('附录 实验复现与文件索引',level=1)
    writer.p('公开代码保存在独立分支 research/semantic-adaptive-v2。先按 README 创建虚拟环境或独立容器，再运行单元测试、配对试验和网络试验。主机和容器路径、镜像来源以及依赖版本均有记录；不得把私有现场图像或数据库一并上传。')
    digests={**analysis['source_sha256'],**{name:row['sha256'] for name,row in analysis['feedback_ablations'].items()}}
    for name,digest in digests.items():
        writer.p(name,False);p=writer.p('SHA256 '+digest,False)
        for run in p.runs: set_font(run,9)
    writer.p('network-pi-fixture-pilot.json 是发现问题时保留的预实验，不进入正式汇总。network-pc.json 和 rd-pc.json 是电脑预实验，也不混入树莓派主表。analysis.json 记录独立复核后的数值和来源摘要。')
    writer.p('论文排版采用通用研究初稿格式，尚未套用学校专用模板。答辩前需要学生核查全部公式、实验记录和引用，补充个人与学校信息，并按导师要求补做真实任务真值与更多重复实验。')
    destination=OUT/'semantic-adaptive-thesis.docx';doc.save(destination)
    print(json.dumps({'docx':str(destination),'tables':len(doc.tables),'equations':8,'source_characters':len(content)},ensure_ascii=False))


if __name__=='__main__': main()
